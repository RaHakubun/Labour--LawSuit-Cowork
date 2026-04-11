#!/usr/bin/env python3
from __future__ import annotations

import argparse
from collections import OrderedDict
from datetime import datetime, timezone
import hashlib
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_SOURCE = "/Users/luobowen/documents/全国地级市平均工资-计算赔偿金封顶用.xlsx"
DEFAULT_TARGET = "utils/city_wage_dataset.py"


def _normalize_text(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _format_wage(v: Any) -> str:
    text = _normalize_text(v)
    if not text:
        raise ValueError("wage cell is empty")
    return text


def _read_rows(source: Path) -> list[tuple[str, str, str]]:
    wb = load_workbook(source, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows: list[tuple[str, str, str]] = []
    current_province = ""
    for row in ws.iter_rows(min_row=2, values_only=True):
        province = _normalize_text(row[0] if len(row) > 0 else "")
        city = _normalize_text(row[1] if len(row) > 1 else "")
        wage = _format_wage(row[2] if len(row) > 2 else "")

        if province:
            current_province = province
        if not current_province:
            raise ValueError(f"province missing before city row: {row!r}")
        if not city:
            city = current_province
        rows.append((current_province, city, wage))
    return rows


def _to_nested(rows: list[tuple[str, str, str]]) -> OrderedDict[str, OrderedDict[str, str]]:
    data: OrderedDict[str, OrderedDict[str, str]] = OrderedDict()
    for province, city, wage in rows:
        if province not in data:
            data[province] = OrderedDict()
        if city in data[province]:
            raise ValueError(f"duplicate city key in dataset: province={province}, city={city}")
        data[province][city] = wage
    return data


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while True:
            chunk = f.read(1024 * 1024)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()


def _render_py(
    source: Path,
    rows: list[tuple[str, str, str]],
    nested: OrderedDict[str, OrderedDict[str, str]],
) -> str:
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    lines: list[str] = []
    lines.append('"""Auto-generated city wage dataset for severance wage cap lookup.')
    lines.append("")
    lines.append("Do not edit manually. Rebuild with:")
    lines.append("  conda run -n Intellectual python scripts/build_city_wage_dataset.py")
    lines.append('"""')
    lines.append("")
    lines.append(f"SOURCE_FILE = {source.name!r}")
    lines.append(f"SOURCE_SHA256 = {_sha256(source)!r}")
    lines.append(f"UPDATED_AT_UTC = {now!r}")
    lines.append(f"RECORD_COUNT = {len(rows)}")
    lines.append(f"PROVINCE_COUNT = {len(nested)}")
    lines.append("")
    lines.append("# province -> city -> latest avg monthly wage (RMB)")
    lines.append("CITY_WAGE_DATA: dict[str, dict[str, str]] = {")
    for province, cities in nested.items():
        lines.append(f"    {province!r}: {{")
        for city, wage in cities.items():
            lines.append(f"        {city!r}: {wage!r},")
        lines.append("    },")
    lines.append("}")
    lines.append("")
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build internal city wage dataset from Excel")
    parser.add_argument("--source", default=DEFAULT_SOURCE)
    parser.add_argument("--target", default=DEFAULT_TARGET)
    args = parser.parse_args()

    source = Path(args.source).expanduser().resolve()
    target = Path(args.target).resolve()
    if not source.exists():
        raise FileNotFoundError(f"source not found: {source}")

    rows = _read_rows(source)
    nested = _to_nested(rows)
    rendered = _render_py(source, rows, nested)

    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(rendered, encoding="utf-8")
    print(f"generated: {target}")
    print(f"records={len(rows)}, provinces={len(nested)}")


if __name__ == "__main__":
    main()

